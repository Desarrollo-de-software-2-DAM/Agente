"""Tarjeta #12: los documentos no deben contradecir las fuentes únicas de verdad."""
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pandas as pd

from scripts.generar_datos_oficiales import generar

DOCS = Path(__file__).resolve().parent.parent / "documents"
CFG = json.loads((DOCS / "config_tienda.json").read_text(encoding="utf-8"))


def _leer(nombre):
    return (DOCS / nombre).read_text(encoding="utf-8")


def test_datos_oficiales_esta_actualizado():
    assert _leer("datos_oficiales.md").replace("\r\n", "\n") == generar().replace("\r\n", "\n"), (
        "Ejecuta: python scripts/generar_datos_oficiales.py"
    )


def test_umbral_envio_gratis_unico():
    umbral = int(CFG["configuracion_envios"]["umbral_envio_gratis"])
    moneda = CFG["configuracion_regional"]["moneda"]
    html = _leer("guia_envios.html")
    assert f"${umbral:,} {moneda}".replace(",", ".") in html
    assert "$50 USD" not in html


def test_html_envios_coincide_con_csv():
    csv = pd.read_csv(DOCS / "precios_envio.csv")
    moneda = CFG["configuracion_regional"]["moneda"]
    html = _leer("guia_envios.html")
    minimos = {
        "Desde $%s %s" % (f"{csv[csv.Zona == 'Local'].Costo_Estandar.min():,.0f}".replace(",", "."), moneda),
        "Desde $%s %s" % (f"{csv[csv.Zona == 'Local'].Costo_Express.min():,.0f}".replace(",", "."), moneda),
        "Desde $%s %s" % (f"{csv[csv.Zona == 'Nacional_Cercano'].Costo_Estandar.min():,.0f}".replace(",", "."), moneda),
        "Desde $%s %s" % (f"{csv[csv.Zona == 'Nacional_Cercano'].Costo_Express.min():,.0f}".replace(",", "."), moneda),
    }
    for texto in minimos:
        assert texto in html


def test_horario_soporte_unico():
    for nombre in ("faq.md", "presentacion_empresa.pptx"):
        if nombre.endswith(".md"):
            assert "24/7" not in _leer(nombre)
    from pptx import Presentation
    texto = " ".join(sh.text_frame.text for s in Presentation(str(DOCS / "presentacion_empresa.pptx")).slides
                     for sh in s.shapes if sh.has_text_frame)
    assert "24/7" not in texto


def test_telefono_unico():
    tel = CFG["tienda"]["telefono"]
    assert tel in _leer("faq.md")


def test_politicas_coinciden_con_config():
    faq = _leer("faq.md")
    assert f"{CFG['politicas']['dias_limite_devolucion']} días naturales" in faq
    moneda = CFG["configuracion_regional"]["moneda"]
    costo = f"{CFG['politicas']['costo_reestocaje_devolucion']:,.0f}".replace(",", ".")
    assert f"${costo} {moneda}" in faq